"""Read and write simple tables: CSV everywhere, XLSX when openpyxl is present.

Accounting data arrives as bank exports, ledger exports and client
spreadsheets. This is the one place that turns any of them into rows
keyed by header, and rows back into a workbook, so every tool reads
files the same way.

openpyxl is optional: without it, .xlsx paths fail with a message that
says what to install, and everything CSV keeps working.
"""
from __future__ import annotations

import csv
import io
import os
from typing import Any, Iterable, Optional

from .books import read_csv_rows


class TableError(ValueError):
    pass


def have_openpyxl() -> bool:
    try:
        import openpyxl  # noqa: F401
        return True
    except ImportError:
        return False


def _openpyxl():
    try:
        import openpyxl
        return openpyxl
    except ImportError:
        raise TableError("reading/writing .xlsx needs the openpyxl package (pip install openpyxl); "
                         "CSV files work without it") from None


def read_table(path: str, *, sheet: Optional[str] = None, header_row: Optional[int] = None,
               max_rows: int = 100000) -> list[dict[str, Any]]:
    ext = os.path.splitext(path)[1].lower()
    if not os.path.isfile(path):
        raise TableError(f"no such file: {path}")
    if ext in (".csv", ".txt", ".tsv"):
        with open(path, encoding="utf-8-sig", errors="replace") as fh:
            return read_csv_rows(fh.read())[:max_rows]
    if ext in (".xlsx", ".xlsm"):
        wb = _openpyxl().load_workbook(path, data_only=True, read_only=True)
        try:
            ws = wb[sheet] if sheet else wb.worksheets[0]
            rows = [list(r) for r in ws.iter_rows(values_only=True)]
        finally:
            wb.close()
        return rows_to_dicts(rows, header_row=header_row)[:max_rows]
    raise TableError(f"unsupported table format {ext or '(none)'}: use .csv or .xlsx")


def rows_to_dicts(rows: list[list[Any]], *, header_row: Optional[int] = None) -> list[dict[str, Any]]:
    rows = [r for r in rows if any(c not in (None, "") for c in r)]
    if not rows:
        return []
    if header_row is None:
        header_row = 0
        for i, r in enumerate(rows[:15]):
            texts = sum(1 for c in r if isinstance(c, str) and c.strip())
            if texts >= max(2, len([c for c in r if c not in (None, '')]) // 2):
                header_row = i
                break
    headers = [str(c).strip() if c not in (None, "") else f"col{j + 1}" for j, c in enumerate(rows[header_row])]
    out = []
    for r in rows[header_row + 1:]:
        out.append({headers[j]: (r[j] if j < len(r) else None) for j in range(len(headers))})
    return out


def write_table(path: str, rows: list[dict[str, Any]], *, sheet: str = "Sheet1",
                columns: Optional[list[str]] = None) -> str:
    columns = columns or _columns(rows)
    os.makedirs(os.path.dirname(os.path.abspath(path)) or ".", exist_ok=True)
    if path.lower().endswith(".csv"):
        with open(path, "w", encoding="utf-8", newline="") as fh:
            w = csv.DictWriter(fh, fieldnames=columns, extrasaction="ignore")
            w.writeheader()
            w.writerows(rows)
        return path
    write_workbook(path, [(sheet, columns, [[r.get(c) for c in columns] for r in rows])])
    return path


def write_workbook(path: str, sheets: Iterable[tuple[str, list[str], list[list[Any]]]], *,
                   money_columns: Iterable[str] = (), title: str = "") -> str:
    """Several sheets, each (name, headers, rows). Values starting with '='
    are written as formulas, so schedules stay live in Excel."""
    op = _openpyxl()
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = op.Workbook()
    wb.remove(wb.active)
    money = set(money_columns)
    for name, headers, rows in sheets:
        ws = wb.create_sheet(_sheet_name(name))
        start = 1
        if title:
            ws.cell(row=1, column=1, value=title).font = Font(bold=True, size=13)
            start = 3
        for j, h in enumerate(headers, 1):
            c = ws.cell(row=start, column=j, value=h)
            c.font = Font(bold=True)
            c.fill = PatternFill("solid", fgColor="DDE6F0")
        for i, row in enumerate(rows, start + 1):
            for j, v in enumerate(row, 1):
                c = ws.cell(row=i, column=j, value=v)
                if headers[j - 1] in money or (isinstance(v, str) and v.startswith("=SUM")):
                    c.number_format = "#,##0.00;[Red](#,##0.00)"
        for j, h in enumerate(headers, 1):
            width = max([len(str(h))] + [len(str(r[j - 1])) for r in rows[:200] if j - 1 < len(r)])
            ws.column_dimensions[get_column_letter(j)].width = min(60, max(10, width + 2))
        ws.freeze_panes = ws.cell(row=start + 1, column=1)
    os.makedirs(os.path.dirname(os.path.abspath(path)) or ".", exist_ok=True)
    wb.save(path)
    return path


def _columns(rows: list[dict[str, Any]]) -> list[str]:
    cols: list[str] = []
    for r in rows:
        for k in r:
            if k not in cols:
                cols.append(k)
    return cols


def _sheet_name(name: str) -> str:
    bad = '[]:*?/\\'
    return "".join("_" if ch in bad else ch for ch in name)[:31] or "Sheet"


def to_csv_text(rows: list[dict[str, Any]]) -> str:
    buf = io.StringIO()
    cols = _columns(rows)
    w = csv.DictWriter(buf, fieldnames=cols)
    w.writeheader()
    w.writerows(rows)
    return buf.getvalue()
