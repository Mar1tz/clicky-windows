"""A brand's marketing desk: the brand kit, its campaigns and every piece in
them - social posts, ads, emails - with their pictures, in one folder.

    <root>/marketing.json     the brand kit: name, voice, audience, colours, fonts, logo, website, hashtags
    <root>/marketing.sqlite   campaigns and pieces
    <root>/Campaigns/<campaign>/   plan.md, images/
    <root>/Exports/<campaign>-<date>/   the kit to hand over or import into a scheduler

What each platform enforces is code here: X's 280 characters, Instagram's
2,200 and 30 hashtags, Google's 30-character headlines, an email's
unsubscribe line. A piece that breaks a rule is still saved - it is a draft -
and the problem is said, so it is fixed before anybody pastes it. Writing
the words and designing the pictures is the model's part; this is where
they are kept and checked. Nothing is posted, sent or paid for from here.
"""
from __future__ import annotations

import csv
import html
import json
import os
import re
import shutil
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from typing import Any, Optional

from extensions.packkit import Store, slug

STATUSES = ("draft", "approved", "scheduled", "posted")
KINDS = ("post", "story", "ad", "email")
CAMPAIGN_STATES = ("planning", "ready", "running", "done")


@dataclass(frozen=True)
class Platform:
    name: str
    limit: int                    # characters a post may have
    size: str                     # the picture's size (agents/tools/design_tools.SIZES)
    fold: int = 0                 # about how much shows before "more"
    hashtags: int = 0             # most hashtags allowed (0: no rule)
    needs_image: bool = False     # a post there is a picture
    links: bool = True            # a link in the text can be clicked


PLATFORMS: dict[str, Platform] = {
    "instagram": Platform("Instagram", 2200, "instagram-portrait", fold=125, hashtags=30, needs_image=True,
                          links=False),
    "facebook": Platform("Facebook", 63206, "facebook-post"),
    "linkedin": Platform("LinkedIn", 3000, "linkedin-post", fold=210),
    "x": Platform("X", 280, "x-post"),
    "threads": Platform("Threads", 500, "instagram-portrait"),
    "tiktok": Platform("TikTok", 4000, "story", needs_image=True, links=False),
    "pinterest": Platform("Pinterest", 500, "pinterest-pin", needs_image=True),
}
ALIASES = {"twitter": "x", "ig": "instagram", "insta": "instagram", "fb": "facebook", "meta": "facebook"}
STORY_SIZE = "story"

# Ads: (field, most characters, at least, at most - how many, hard rule?)
AD_RULES: dict[str, dict[str, Any]] = {
    "google-search": {"name": "Google search ad", "headlines": (30, 3, 15), "descriptions": (90, 2, 4),
                      "hard": True},
    "meta": {"name": "Meta (Facebook/Instagram) ad", "primary_text": 125, "headline": 40, "description": 30,
             "hard": False, "size": "instagram-post"},
    "linkedin": {"name": "LinkedIn ad", "primary_text": 150, "headline": 70, "hard": False,
                 "size": "linkedin-post"},
}
AD_ALIASES = {"google": "google-search", "google-ads": "google-search", "facebook": "meta", "instagram": "meta",
              "meta-ads": "meta", "linkedin-ads": "linkedin"}

SUBJECT_MAX = 60                  # longer is cut off in most inboxes (about 40 on a phone)
PREVIEW_MAX = 140
SPAMMY = re.compile(r"\b(?:100% free|act now|buy now!|click here|guaranteed|risk[- ]free|winner|cash bonus|"
                    r"no obligation)\b", re.I)
_URL = re.compile(r"https?://\S+|www\.\S+", re.I)
_TAG = re.compile(r"(?<![\w&])#(\w+)")


def _now() -> str:
    return datetime.now().isoformat(timespec="seconds")


def _date(value: Any) -> str:
    if not value:
        return ""
    try:
        return date.fromisoformat(str(value).strip()[:10]).isoformat()
    except ValueError:
        raise ValueError(f"{value!r} is not a date (YYYY-MM-DD)") from None


def _time(value: Any) -> str:
    if not value:
        return ""
    m = re.match(r"^\s*(\d{1,2})[:.](\d{2})\s*$", str(value))
    if not m or int(m.group(1)) > 23 or int(m.group(2)) > 59:
        raise ValueError(f"{value!r} is not a time (HH:MM, 24-hour)")
    return f"{int(m.group(1)):02d}:{m.group(2)}"


def platform_key(name: str) -> str:
    key = re.sub(r"[^a-z]", "", (name or "").lower())
    key = ALIASES.get(key, key)
    if key not in PLATFORMS:
        raise ValueError(f"{name!r} is not a platform here - one of: {', '.join(PLATFORMS)}")
    return key


def ad_key(name: str) -> str:
    key = re.sub(r"[^a-z-]", "", (name or "").lower().replace(" ", "-"))
    key = AD_ALIASES.get(key, key)
    if key not in AD_RULES:
        raise ValueError(f"{name!r} is not an ad type here - one of: {', '.join(AD_RULES)}")
    return key


def x_length(text: str) -> int:
    """X counts every link as 23 characters, whatever its length."""
    return len(_URL.sub("x" * 23, text or ""))


def hashtags_in(text: str, extra: list[str]) -> list[str]:
    tags = [t.lower() for t in _TAG.findall(text or "")]
    tags += [str(t).lstrip("#").lower() for t in extra or [] if str(t).strip("# ")]
    return list(dict.fromkeys(tags))


def _problem(title: str, action: str, severity: str = "fix") -> dict[str, str]:
    return {"title": title, "action": action, "severity": severity}


def post_problems(platform: str, text: str, hashtags: list[str], kind: str = "post") -> list[dict[str, str]]:
    """What stops a post going out as written, on that platform."""
    p = PLATFORMS[platform]
    full = (text or "") + ("\n\n" + " ".join("#" + t for t in hashtags) if hashtags else "")
    out = []
    length = x_length(full) if platform == "x" else len(full)
    if not (text or "").strip() and kind != "story":
        out.append(_problem(f"{p.name}: the post has no text", "write it"))
    if length > p.limit:
        out.append(_problem(f"{p.name}: {length} characters - the limit is {p.limit}",
                            f"cut {length - p.limit} characters"))
    tags = hashtags_in(full, [])
    if p.hashtags and len(tags) > p.hashtags:
        out.append(_problem(f"{p.name}: {len(tags)} hashtags - at most {p.hashtags}",
                            f"drop {len(tags) - p.hashtags}"))
    if not p.links and _URL.search(text or ""):
        out.append(_problem(f"{p.name}: links in the text cannot be clicked there",
                            "say 'link in bio' and put the link in the profile", "review"))
    return out


def ad_problems(kind: str, data: dict[str, Any]) -> list[dict[str, str]]:
    rules = AD_RULES[kind]
    out = []
    sev = "fix" if rules["hard"] else "review"
    verb = "the limit is" if rules["hard"] else "more is cut off - keep it to"
    for field_ in ("headlines", "descriptions"):
        if field_ not in rules:
            continue
        most, least, max_count = rules[field_]
        items = [str(x) for x in data.get(field_) or []]
        if len(items) < least:
            out.append(_problem(f"{rules['name']}: {len(items)} {field_} - it needs at least {least}",
                                f"write {least - len(items)} more"))
        if len(items) > max_count:
            out.append(_problem(f"{rules['name']}: {len(items)} {field_} - at most {max_count}",
                                f"drop {len(items) - max_count}"))
        for item in items:
            if len(item) > most:
                out.append(_problem(f"{rules['name']}: '{item[:40]}' is {len(item)} characters - {verb} {most}",
                                    "shorten it", sev))
    for field_ in ("primary_text", "headline", "description"):
        if field_ in rules and isinstance(rules[field_], int):
            value = str(data.get(field_) or "")
            if len(value) > rules[field_]:
                out.append(_problem(f"{rules['name']}: the {field_.replace('_', ' ')} is {len(value)} characters "
                                    f"- {verb} {rules[field_]}", "shorten it", sev))
    if kind == "meta" and not str(data.get("primary_text") or "").strip():
        out.append(_problem("Meta ad: no primary text", "write it"))
    if kind == "google-search" and not str(data.get("link") or "").strip():
        out.append(_problem("Google search ad: no final URL (the page it opens)", "add the link"))
    return out


def email_problems(subject: str, preview: str, body: str) -> list[dict[str, str]]:
    out = []
    if not (subject or "").strip():
        out.append(_problem("Email: no subject line", "write it"))
    elif len(subject) > SUBJECT_MAX:
        out.append(_problem(f"Email: the subject is {len(subject)} characters - most inboxes cut it off after "
                            f"about {SUBJECT_MAX}", "shorten it", "review"))
    letters = [c for c in subject or "" if c.isalpha()]
    if len(letters) >= 6 and all(c.isupper() for c in letters):
        out.append(_problem("Email: the subject is all capitals - it reads as shouting and trips spam filters",
                            "write it in normal case"))
    if (subject or "").count("!") > 1 or SPAMMY.search(subject or ""):
        out.append(_problem("Email: the subject looks like spam (several '!' or words spam filters watch)",
                            "say it plainly", "review"))
    if not (preview or "").strip():
        out.append(_problem("Email: no preview text - the inbox shows the first line of the body instead",
                            "add one line that adds to the subject", "review"))
    elif len(preview) > PREVIEW_MAX:
        out.append(_problem(f"Email: the preview text is {len(preview)} characters - inboxes show about "
                            f"{PREVIEW_MAX}", "shorten it", "review"))
    if not (body or "").strip():
        out.append(_problem("Email: no body", "write it"))
    elif not re.search(r"unsubscribe", body, re.I):
        out.append(_problem("Email: no unsubscribe line - a marketing email needs one, and the sender's postal "
                            "address, in most countries", "add them at the foot"))
    return out


class Desk(Store):
    SETTINGS_FILE = "marketing.json"
    DB = "marketing.sqlite"
    FOLDERS = ("Campaigns", "Exports", "Brand")
    SCHEMA = """
    CREATE TABLE IF NOT EXISTS campaigns (
        id INTEGER PRIMARY KEY, slug TEXT NOT NULL UNIQUE, name TEXT NOT NULL, goal TEXT NOT NULL DEFAULT '',
        audience TEXT NOT NULL DEFAULT '', offer TEXT NOT NULL DEFAULT '', channels TEXT NOT NULL DEFAULT '[]',
        start TEXT NOT NULL DEFAULT '', end TEXT NOT NULL DEFAULT '', budget REAL NOT NULL DEFAULT 0,
        status TEXT NOT NULL DEFAULT 'planning', plan TEXT NOT NULL DEFAULT '', created TEXT NOT NULL
    );
    CREATE TABLE IF NOT EXISTS pieces (
        id INTEGER PRIMARY KEY, campaign_id INTEGER NOT NULL, kind TEXT NOT NULL, platform TEXT NOT NULL,
        day TEXT NOT NULL DEFAULT '', time TEXT NOT NULL DEFAULT '', text TEXT NOT NULL DEFAULT '',
        hashtags TEXT NOT NULL DEFAULT '[]', link TEXT NOT NULL DEFAULT '', image TEXT NOT NULL DEFAULT '',
        data TEXT NOT NULL DEFAULT '{}', status TEXT NOT NULL DEFAULT 'draft', created TEXT NOT NULL
    );
    """

    # -- the brand kit -------------------------------------------------------------------------------------------

    BRAND_FIELDS = ("name", "tagline", "voice", "audience", "colours", "fonts", "logo", "website", "hashtags",
                    "currency")

    def brand(self) -> dict[str, Any]:
        b = {k: self.cfg.get(k) for k in self.BRAND_FIELDS}
        b["name"] = b["name"] or "My brand"
        for k in ("colours", "fonts", "hashtags"):
            b[k] = list(b.get(k) or [])
        if b["logo"]:
            b["logo_path"] = os.path.join(self.root, b["logo"])
        return b

    def set_brand(self, **fields: Any) -> dict[str, Any]:
        for k, v in fields.items():
            if k not in self.BRAND_FIELDS or v is None:
                continue
            if k == "colours":
                v = [str(c).strip() for c in (v if isinstance(v, list) else str(v).split(","))]
                bad = [c for c in v if not re.fullmatch(r"#[0-9a-fA-F]{6}", c)]
                if bad:
                    raise ValueError(f"colours are hex like #14213D - not {', '.join(bad)}")
            elif k in ("fonts", "hashtags"):
                v = [str(x).strip().lstrip("#" if k == "hashtags" else "") for x in
                     (v if isinstance(v, list) else str(v).split(",")) if str(x).strip()]
            elif k == "logo" and v:
                v = self._inside(str(v))
            self.cfg[k] = v
        with self.lock, open(os.path.join(self.root, self.SETTINGS_FILE), "w", encoding="utf-8") as fh:
            json.dump(self.cfg, fh, indent=2)
        return self.brand()

    def _inside(self, path: str) -> str:
        """A file of the brand's, as a path relative to the folder; it must be in it."""
        full = os.path.normpath(path if os.path.isabs(path) else os.path.join(self.root, path))
        if os.path.commonpath([os.path.normcase(full), os.path.normcase(self.root)]) != os.path.normcase(self.root):
            raise ValueError(f"{path} is outside the marketing folder - copy it into {self.folder('Brand')}")
        if not os.path.isfile(full):
            raise ValueError(f"there is no file {path}")
        return os.path.relpath(full, self.root)

    # -- campaigns ----------------------------------------------------------------------------------------------

    def _campaign_row(self, r: dict[str, Any]) -> dict[str, Any]:
        r = dict(r)
        r["channels"] = json.loads(r["channels"] or "[]")
        r["folder"] = os.path.join(self.root, "Campaigns", r["slug"])
        return r

    def campaigns(self, status: str = "") -> list[dict[str, Any]]:
        sql = "SELECT * FROM campaigns" + (" WHERE status = ?" if status else "") + " ORDER BY id DESC"
        return [self._campaign_row(r) for r in self.rows(sql, (status,) if status else ())]

    def campaign(self, ref: Any) -> dict[str, Any]:
        """By id, its name or the start of it, or "latest"."""
        text = str(ref or "").strip()
        if not text or text.lower() in ("latest", "last", "current", "newest"):
            r = self.one("SELECT * FROM campaigns ORDER BY id DESC LIMIT 1")
            if r is None:
                raise LookupError("there is no campaign yet - create one first")
            return self._campaign_row(r)
        if text.isdigit():
            r = self.one("SELECT * FROM campaigns WHERE id = ?", (int(text),))
        else:
            r = self.one("SELECT * FROM campaigns WHERE slug = ?", (slug(text),)) or self.one(
                "SELECT * FROM campaigns WHERE lower(name) LIKE ? ORDER BY id DESC", (text.lower() + "%",))
        if r is None:
            raise LookupError(f"no campaign {text!r}")
        return self._campaign_row(r)

    def add_campaign(self, name: str, *, goal: str = "", audience: str = "", offer: str = "",
                     channels: Optional[list[str]] = None, start: str = "", end: str = "",
                     budget: float = 0) -> dict[str, Any]:
        name = (name or "").strip()
        if not name:
            raise ValueError("the campaign needs a name")
        start = _date(start) or date.today().isoformat()
        end = _date(end) or (date.fromisoformat(start) + timedelta(days=13)).isoformat()
        if end < start:
            raise ValueError(f"the campaign ends ({end}) before it starts ({start})")
        chans = []
        for c in channels or []:
            c = str(c).strip().lower()
            if not c:
                continue
            if c in ("email", "newsletter"):
                chans.append("email")
                continue
            try:
                chans.append(platform_key(c))
                continue
            except ValueError:
                pass
            try:                                   # "google ads", "meta-ads", "linkedin ads"
                chans.append(ad_key(re.sub(r"[\s-]*ads?$", "", c)) + "-ads")
            except ValueError:
                raise ValueError(f"{c!r} is not a channel here - a platform ({', '.join(PLATFORMS)}), email, "
                                 f"or ads ({', '.join(a + ' ads' for a in AD_RULES)})") from None
        base, n = slug(name), 1
        s = base
        while self.one("SELECT id FROM campaigns WHERE slug = ?", (s,)):
            n += 1
            s = f"{base}-{n}"
        cid = self.write("INSERT INTO campaigns (slug, name, goal, audience, offer, channels, start, end, budget, "
                         "created) VALUES (?,?,?,?,?,?,?,?,?,?)",
                         (s, name, goal or "", audience or "", offer or "", json.dumps(list(dict.fromkeys(chans))),
                          start, end, float(budget or 0), _now()))
        os.makedirs(os.path.join(self.root, "Campaigns", s, "images"), exist_ok=True)
        return self.campaign(cid)

    def save_plan(self, ref: Any, text: str) -> dict[str, Any]:
        c = self.campaign(ref)
        if not (text or "").strip():
            raise ValueError("the plan is empty")
        path = os.path.join(c["folder"], "plan.md")
        os.makedirs(c["folder"], exist_ok=True)
        with open(path, "w", encoding="utf-8") as fh:
            fh.write(text.strip() + "\n")
        self.write("UPDATE campaigns SET plan = ? WHERE id = ?", (path, c["id"]))
        return {**self.campaign(c["id"]), "words": len(text.split())}

    def set_campaign_status(self, ref: Any, status: str) -> dict[str, Any]:
        if status not in CAMPAIGN_STATES:
            raise ValueError(f"status is one of {', '.join(CAMPAIGN_STATES)}")
        c = self.campaign(ref)
        self.write("UPDATE campaigns SET status = ? WHERE id = ?", (status, c["id"]))
        return self.campaign(c["id"])

    # -- pieces -------------------------------------------------------------------------------------------------

    def _piece_row(self, r: dict[str, Any]) -> dict[str, Any]:
        r = dict(r)
        r["hashtags"] = json.loads(r["hashtags"] or "[]")
        r["data"] = json.loads(r["data"] or "{}")
        r["image_path"] = os.path.join(self.root, r["image"]) if r["image"] else ""
        r["problems"] = self.problems(r)
        return r

    def piece(self, piece_id: Any) -> dict[str, Any]:
        r = self.one("SELECT * FROM pieces WHERE id = ?", (int(piece_id),))
        if r is None:
            raise LookupError(f"no piece #{piece_id}")
        return self._piece_row(r)

    def pieces(self, campaign: Any = None, *, kind: str = "", platform: str = "", status: str = "") -> list[dict]:
        sql, params = "SELECT * FROM pieces WHERE 1=1", []
        if campaign not in (None, ""):
            sql += " AND campaign_id = ?"
            params.append(self.campaign(campaign)["id"])
        for col, val in (("kind", kind), ("platform", platform), ("status", status)):
            if val:
                sql += f" AND {col} = ?"
                params.append(val)
        return [self._piece_row(r) for r in self.rows(sql + " ORDER BY day, time, id", params)]

    def problems(self, p: dict[str, Any]) -> list[dict[str, str]]:
        data = p["data"] if isinstance(p["data"], dict) else json.loads(p["data"] or "{}")
        if p["kind"] == "email":
            return email_problems(data.get("subject", ""), data.get("preview", ""), p["text"])
        if p["kind"] == "ad":
            return ad_problems(p["platform"], {**data, "link": p["link"]})
        tags = p["hashtags"] if isinstance(p["hashtags"], list) else json.loads(p["hashtags"] or "[]")
        return post_problems(p["platform"], p["text"], tags, p["kind"])

    def _insert(self, campaign: dict[str, Any], kind: str, platform: str, *, day: str = "", time: str = "",
                text: str = "", hashtags: Optional[list[str]] = None, link: str = "",
                data: Optional[dict[str, Any]] = None) -> dict[str, Any]:
        pid = self.write("INSERT INTO pieces (campaign_id, kind, platform, day, time, text, hashtags, link, data, "
                         "created) VALUES (?,?,?,?,?,?,?,?,?,?)",
                         (campaign["id"], kind, platform, _date(day), _time(time), text or "",
                          json.dumps([str(t).lstrip("#") for t in hashtags or [] if str(t).strip("# ")]),
                          link or "", json.dumps(data or {}), _now()))
        return self.piece(pid)

    def add_post(self, campaign: Any, platform: str, text: str, *, day: str = "", time: str = "",
                 hashtags: Optional[list[str]] = None, link: str = "", kind: str = "post",
                 cta: str = "") -> dict[str, Any]:
        if kind not in ("post", "story"):
            raise ValueError("a social piece is a post or a story")
        c = self.campaign(campaign)
        return self._insert(c, kind, platform_key(platform), day=day, time=time, text=text, hashtags=hashtags,
                            link=link, data={"cta": cta} if cta else {})

    def add_ad(self, campaign: Any, platform: str, *, headlines: Optional[list[str]] = None,
               descriptions: Optional[list[str]] = None, primary_text: str = "", headline: str = "",
               description: str = "", link: str = "", cta: str = "", day: str = "") -> dict[str, Any]:
        c = self.campaign(campaign)
        data = {k: v for k, v in (("headlines", list(headlines or [])), ("descriptions", list(descriptions or [])),
                                  ("primary_text", primary_text), ("headline", headline),
                                  ("description", description), ("cta", cta)) if v}
        return self._insert(c, "ad", ad_key(platform), day=day, text=primary_text or "\n".join(headlines or []),
                            link=link, data=data)

    def add_email(self, campaign: Any, subject: str, body: str, *, preview: str = "", day: str = "",
                  time: str = "", audience: str = "", link: str = "", cta: str = "") -> dict[str, Any]:
        c = self.campaign(campaign)
        data = {k: v for k, v in (("subject", subject), ("preview", preview), ("audience", audience),
                                  ("cta", cta)) if v}
        return self._insert(c, "email", "email", day=day, time=time, text=body, link=link, data=data)

    def update_piece(self, piece_id: Any, **fields: Any) -> dict[str, Any]:
        p = self.piece(piece_id)
        cols, params = [], []
        data = dict(p["data"])
        for k, v in fields.items():
            if v is None:
                continue
            if k == "status":
                if v not in STATUSES:
                    raise ValueError(f"status is one of {', '.join(STATUSES)}")
                cols.append("status = ?"); params.append(v)
            elif k in ("day", "date"):
                cols.append("day = ?"); params.append(_date(v))
            elif k == "time":
                cols.append("time = ?"); params.append(_time(v))
            elif k in ("text", "link"):
                cols.append(f"{k} = ?"); params.append(str(v))
            elif k == "hashtags":
                cols.append("hashtags = ?")
                params.append(json.dumps([str(t).lstrip("#") for t in v if str(t).strip("# ")]))
            elif k in ("subject", "preview", "headlines", "descriptions", "primary_text", "headline",
                       "description", "cta", "audience"):
                data[k] = v
        cols.append("data = ?"); params.append(json.dumps(data))
        self.write(f"UPDATE pieces SET {', '.join(cols)} WHERE id = ?", (*params, p["id"]))
        return self.piece(p["id"])

    def image_target(self, piece_id: Any = None, name: str = "", campaign: Any = None,
                     size: str = "") -> tuple[str, str, Optional[dict]]:
        """(where the picture goes, its size, the piece) for a piece or a free-standing picture."""
        if piece_id not in (None, ""):
            p = self.piece(piece_id)
            c = self.campaign(p["campaign_id"])
            if p["kind"] == "email":
                default = "email-header"
            elif p["kind"] == "ad":
                default = AD_RULES[p["platform"]].get("size", "")
                if not default:
                    raise ValueError("a Google search ad has no picture")
            else:
                default = STORY_SIZE if p["kind"] == "story" else PLATFORMS[p["platform"]].size
            return (os.path.join(c["folder"], "images", f"{p['id']}-{p['platform']}-{p['kind']}.png"),
                    size or default, p)
        if not name:
            raise ValueError("say which piece the picture is for, or give it a name")
        try:                                       # a campaign's picture - or, with none, a design of its own
            folder = os.path.join(self.campaign(campaign)["folder"], "images")
        except LookupError:
            folder = os.path.join(self.root, "Designs")
        return os.path.join(folder, f"{slug(name)}.png"), size or "instagram-post", None

    def set_image(self, piece_id: Any, path: str) -> dict[str, Any]:
        rel = os.path.relpath(path, self.root)
        self.write("UPDATE pieces SET image = ? WHERE id = ?", (rel, int(piece_id)))
        return self.piece(piece_id)

    # -- the checks -----------------------------------------------------------------------------------------------

    def check(self, ref: Any) -> dict[str, Any]:
        c = self.campaign(ref)
        pieces = self.pieces(c["id"])
        issues: list[dict[str, str]] = []
        for p in pieces:
            for prob in p["problems"]:
                issues.append({**prob, "piece": p["id"]})
            if p["kind"] in ("post", "story") and not p["image"] and (
                    PLATFORMS[p["platform"]].needs_image or p["kind"] == "story"):
                issues.append({**_problem(f"#{p['id']} {PLATFORMS[p['platform']].name}: needs a picture",
                                          "make one with make_image"), "piece": p["id"]})
            if p["day"] and not (c["start"] <= p["day"] <= c["end"]):
                issues.append({**_problem(f"#{p['id']} is dated {p['day']}, outside the campaign "
                                          f"({c['start']} to {c['end']})", "move it, or the campaign's dates",
                                          "review"), "piece": p["id"]})
            if p["kind"] in ("post", "story") and not p["day"]:
                issues.append({**_problem(f"#{p['id']} {PLATFORMS[p['platform']].name} has no date",
                                          "give it a day"), "piece": p["id"]})
        slots: dict[tuple, list[int]] = {}
        for p in pieces:
            if p["kind"] in ("post", "story") and p["day"]:
                slots.setdefault((p["platform"], p["day"], p["time"]), []).append(p["id"])
        for (plat, day, time), ids in slots.items():
            if len(ids) > 1:
                issues.append(_problem(f"{PLATFORMS[plat].name}: {len(ids)} posts at {day} {time or '(no time)'} "
                                       f"(#{', #'.join(map(str, ids))})", "space them out", "review"))
        used = {p["platform"] for p in pieces if p["kind"] in ("post", "story")}
        kinds = {p["kind"] for p in pieces}
        for ch in c["channels"]:
            if ch in PLATFORMS and ch not in used:
                issues.append(_problem(f"{PLATFORMS[ch].name} is a channel of this campaign with no posts",
                                       "write some, or drop the channel", "review"))
            if ch == "email" and "email" not in kinds:
                issues.append(_problem("email is a channel of this campaign with no email", "write one", "review"))
            if ch.endswith("-ads") and not any(p["kind"] == "ad" and p["platform"] == ch[:-4] for p in pieces):
                issues.append(_problem(f"{AD_RULES[ch[:-4]]['name']}s are a channel with no ad", "write them",
                                       "review"))
        if not c["plan"]:
            issues.append(_problem("the campaign has no plan", "write it with save_campaign_plan", "review"))
        return {"campaign": c, "pieces": len(pieces), "issues": issues,
                "fix": sum(1 for i in issues if i["severity"] == "fix")}

    def overview(self) -> dict[str, Any]:
        today = date.today().isoformat()
        week = (date.today() + timedelta(days=7)).isoformat()
        camps = self.campaigns()
        pieces = self.pieces()
        upcoming = [p for p in pieces if p["kind"] in ("post", "story") and p["day"] and today <= p["day"] <= week]
        return {"brand": self.brand(), "campaigns": camps,
                "live": [c for c in camps if c["start"] <= today <= c["end"] and c["status"] != "done"],
                "counts": {k: sum(1 for p in pieces if p["kind"] == k) for k in KINDS},
                "images": sum(1 for p in pieces if p["image"]),
                "upcoming": upcoming,
                "problems": sum(len(p["problems"]) for p in pieces)}

    # -- the kit ----------------------------------------------------------------------------------------------------

    def export(self, ref: Any) -> dict[str, Any]:
        """The campaign as a folder to hand over: the plan, every picture, a
        posting schedule (CSV), the same in Hootsuite's bulk-upload layout,
        the ads, each email as HTML and text, and a calendar workbook."""
        c = self.campaign(ref)
        pieces = self.pieces(c["id"])
        out = os.path.join(self.root, "Exports", f"{c['slug']}-{date.today().isoformat()}")
        if os.path.isdir(out):
            shutil.rmtree(out)
        os.makedirs(os.path.join(out, "images"))
        files = []
        if c["plan"] and os.path.isfile(c["plan"]):
            shutil.copyfile(c["plan"], os.path.join(out, "plan.md"))
            files.append("plan.md")

        def picture(p):
            if not p["image_path"] or not os.path.isfile(p["image_path"]):
                return ""
            name = os.path.basename(p["image_path"])
            shutil.copyfile(p["image_path"], os.path.join(out, "images", name))
            return f"images/{name}"

        social = [p for p in pieces if p["kind"] in ("post", "story")]
        with open(os.path.join(out, "schedule.csv"), "w", encoding="utf-8-sig", newline="") as fh:
            w = csv.writer(fh)
            w.writerow(["date", "time", "platform", "kind", "text", "hashtags", "link", "image", "status"])
            for p in social:
                w.writerow([p["day"], p["time"], PLATFORMS[p["platform"]].name, p["kind"], p["text"],
                            " ".join("#" + t for t in p["hashtags"]), p["link"], picture(p), p["status"]])
        files.append("schedule.csv")
        # Hootsuite's bulk composer layout: no header; the date as
        # "DD/MM/YYYY HH:mm" on a five-minute mark, the message, a link.
        with open(os.path.join(out, "hootsuite.csv"), "w", encoding="utf-8", newline="") as fh:
            w = csv.writer(fh)
            for p in social:
                if not p["day"]:
                    continue
                hh, mm = (p["time"] or "09:00").split(":")
                mm = f"{int(mm) // 5 * 5:02d}"
                y, mo, d = p["day"].split("-")
                tags = " ".join("#" + t for t in p["hashtags"])
                w.writerow([f"{d}/{mo}/{y} {hh}:{mm}", (p["text"] + ("\n\n" + tags if tags else "")).strip(),
                            p["link"]])
        files.append("hootsuite.csv")
        ads = [p for p in pieces if p["kind"] == "ad"]
        if ads:
            with open(os.path.join(out, "ads.csv"), "w", encoding="utf-8-sig", newline="") as fh:
                w = csv.writer(fh)
                w.writerow(["type", "headlines", "descriptions", "primary text", "headline", "description", "link",
                            "call to action", "image"])
                for p in ads:
                    d = p["data"]
                    w.writerow([AD_RULES[p["platform"]]["name"], " | ".join(d.get("headlines", [])),
                                " | ".join(d.get("descriptions", [])), d.get("primary_text", ""),
                                d.get("headline", ""), d.get("description", ""), p["link"], d.get("cta", ""),
                                picture(p)])
            files.append("ads.csv")
        emails = [p for p in pieces if p["kind"] == "email"]
        for i, p in enumerate(emails, 1):
            base = f"email-{i}-{slug(p['data'].get('subject', 'email'))[:40]}"
            img = picture(p)
            with open(os.path.join(out, base + ".html"), "w", encoding="utf-8") as fh:
                fh.write(email_html(self.brand(), p, img))
            with open(os.path.join(out, base + ".txt"), "w", encoding="utf-8") as fh:
                fh.write(f"Subject: {p['data'].get('subject', '')}\nPreview: {p['data'].get('preview', '')}\n"
                         f"Send: {p['day']} {p['time']}\n\n{p['text']}\n")
            files += [base + ".html", base + ".txt"]
        calendar = self._calendar(out, c, pieces)
        if calendar:
            files.append(calendar)
        images = sorted(os.listdir(os.path.join(out, "images")))
        return {"campaign": c, "folder": out, "files": files, "images": len(images), "posts": len(social),
                "ads": len(ads), "emails": len(emails), "issues": self.check(c["id"])["issues"]}

    def _calendar(self, out: str, c: dict[str, Any], pieces: list[dict[str, Any]]) -> str:
        try:
            import openpyxl
            from openpyxl.styles import Alignment, Font, PatternFill
        except ImportError:
            return ""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Calendar"
        ws.append([f"{c['name']} - {c['start']} to {c['end']}"])
        ws["A1"].font = Font(bold=True, size=14)
        ws.append([])
        ws.append(["Date", "Time", "Channel", "Kind", "Text / subject", "Picture", "Status", "Needs fixing"])
        head = PatternFill("solid", fgColor="14213D")
        for cell in ws[3]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = head
        for p in sorted(pieces, key=lambda p: (p["day"] or "9999", p["time"], p["id"])):
            channel = (PLATFORMS[p["platform"]].name if p["platform"] in PLATFORMS
                       else AD_RULES.get(p["platform"], {}).get("name", "Email"))
            words = p["data"].get("subject") if p["kind"] == "email" else p["text"]
            ws.append([p["day"], p["time"], channel, p["kind"], words, os.path.basename(p["image"]) if p["image"]
                       else "", p["status"], "; ".join(x["title"] for x in p["problems"])])
        for col, width in zip("ABCDEFGH", (12, 7, 22, 8, 70, 28, 10, 40)):
            ws.column_dimensions[col].width = width
        for row in ws.iter_rows(min_row=4):
            row[4].alignment = Alignment(wrap_text=True, vertical="top")
        ws.freeze_panes = "A4"
        wb.save(os.path.join(out, "calendar.xlsx"))
        return "calendar.xlsx"


def email_html(brand: dict[str, Any], p: dict[str, Any], image: str = "") -> str:
    """A plain, safe HTML email: the brand's colour, the picture, the body's
    paragraphs, one button. Every word is escaped."""
    colours = brand.get("colours") or ["#14213D", "#FCA311"]
    main, accent = colours[0], colours[1] if len(colours) > 1 else colours[0]
    font = (brand.get("fonts") or ["Segoe UI"])[0]
    paras = "".join(f'<p style="margin:0 0 16px">{html.escape(par).replace(chr(10), "<br>")}</p>'
                    for par in re.split(r"\n\s*\n", p["text"].strip()) if par.strip())
    pic = f'<img src="{html.escape(image)}" width="600" style="width:100%;display:block" alt="">' if image else ""
    cta = p["data"].get("cta") or ("Find out more" if p["link"] else "")
    button = (f'<p style="margin:24px 0"><a href="{html.escape(p["link"])}" style="background:{accent};color:{main};'
              f'padding:14px 24px;border-radius:6px;text-decoration:none;font-weight:bold">{html.escape(cta)}</a></p>'
              if p["link"] else "")
    return (f'<!doctype html><html><head><meta charset="utf-8"><title>{html.escape(p["data"].get("subject", ""))}'
            f'</title></head><body style="margin:0;background:#f4f4f4">'
            f'<span style="display:none">{html.escape(p["data"].get("preview", ""))}</span>'
            f'<table width="100%" cellpadding="0" cellspacing="0"><tr><td align="center">'
            f'<table width="600" cellpadding="0" cellspacing="0" style="background:#fff;font-family:\'{html.escape(font)}\','
            f'Arial,sans-serif;font-size:16px;line-height:1.5;color:#222">'
            f'<tr><td style="background:{main};color:#fff;padding:20px 32px;font-size:20px;font-weight:bold">'
            f'{html.escape(brand.get("name") or "")}</td></tr><tr><td>{pic}</td></tr>'
            f'<tr><td style="padding:32px">{paras}{button}</td></tr></table></td></tr></table></body></html>')
